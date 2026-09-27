"""Synthetic country eligibility, ordinal inference and measurement-screen safeguards."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from mp1.cross_cultural import (
    ITEMS,
    MODELS,
    SUBDIMENSIONS,
    audit_external_sample,
    inspect_workbook,
    spearman_evidence,
)


@pytest.fixture
def country_data() -> pd.DataFrame:
    rng = np.random.default_rng(74)
    frame = pd.DataFrame(rng.integers(1, 6, (42, len(ITEMS))), columns=ITEMS).astype(float)
    frame["RespondentID"] = [f"synthetic_{i}" for i in range(42)]
    frame["Country"] = np.repeat(["USA", "Indonesia"], 21)
    frame["AnalyticSample"] = "Yes"
    frame["GenAI_Frequency"] = "Sometimes"
    frame.loc[[20, 41], ITEMS] = np.nan
    frame.loc[[20, 41], "AnalyticSample"] = "No"
    frame.loc[[20, 41], "GenAI_Frequency"] = "Rarely"
    return frame


def test_country_identities_and_source_eligibility_are_preserved(country_data) -> None:
    original = country_data.copy(deep=True)
    analytic, accounting = audit_external_sample(country_data, ["USA", "Indonesia"])
    assert analytic.Country.value_counts().to_dict() == {"USA": 20, "Indonesia": 20}
    assert accounting.analytic_records.tolist() == [20, 20]
    assert accounting.whole_battery_missing_nonactive.tolist() == [1, 1]
    assert accounting.source_rows_deleted.eq(0).all()
    pd.testing.assert_frame_equal(country_data, original)


@pytest.mark.parametrize(
    "corruption", ["country", "identifier", "noninteger", "partial", "flag", "frequency"]
)
def test_unreviewed_country_or_item_changes_fail_closed(country_data, corruption) -> None:
    if corruption == "country":
        country_data.loc[0, "Country"] = "US"
    elif corruption == "identifier":
        country_data.loc[0, "RespondentID"] = country_data.loc[1, "RespondentID"]
    elif corruption == "noninteger":
        country_data.loc[0, "GAID1"] = 1.5
    elif corruption == "partial":
        country_data.loc[0, "GAID1"] = np.nan
    elif corruption == "flag":
        country_data.loc[20, "AnalyticSample"] = "Yes"
    else:
        country_data.loc[0, "GenAI_Frequency"] = "Unknown"
    with pytest.raises(ValueError):
        audit_external_sample(country_data, ["USA", "Indonesia"])


def synthetic_workbook(path: Path, frame: pd.DataFrame) -> None:
    """Create a small synthetic fixture; no observed source records are embedded."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Survey Data"
    sheet.append(list(frame))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else value for value in row])
    codebook = book.create_sheet("Codebook")
    codebook.append(
        ["Variable", "Question", "Construct", "Subdimension", "Response coding / Values", "Notes"]
    )
    membership = {
        item: (construct, SUBDIMENSIONS[name])
        for construct, groups in MODELS.items()
        for name, items in groups.items()
        for item in items
    }
    for column in frame:
        construct, dimension = membership.get(column, ("Metadata", "Metadata"))
        codebook.append(
            [
                column,
                "Synthetic description",
                construct,
                dimension,
                "1 = Strongly disagree to 5 = Strongly agree",
                "Synthetic fixture",
            ]
        )
    book.save(path)
    book.close()


def test_entire_workbook_is_inspected_without_modification(tmp_path, country_data) -> None:
    path = tmp_path / "synthetic.xlsx"
    synthetic_workbook(path, country_data)
    before = path.read_bytes()
    data, codebook, inventory = inspect_workbook(path)
    assert len(data) == 42
    assert codebook.Variable.tolist() == list(data)
    assert inventory.all_cells_inspected.all()
    assert inventory.formula_cells.sum() == 0
    assert path.read_bytes() == before
    workbook = openpyxl.load_workbook(path)
    workbook["Codebook"]["B2"] = "=1+1"
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="Formula"):
        inspect_workbook(path)


def test_missing_codebook_field_prevents_variable_selection(tmp_path, country_data) -> None:
    path = tmp_path / "synthetic.xlsx"
    synthetic_workbook(path, country_data)
    workbook = openpyxl.load_workbook(path)
    workbook["Codebook"].delete_rows(2)
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="cover every data column"):
        inspect_workbook(path)


def test_spearman_ties_monotonic_invariance_and_repeatability() -> None:
    x = np.tile(np.arange(1, 6), 20)
    y = np.roll(x, 1)
    first = spearman_evidence(x, y, 499, 500, 19)
    assert first == spearman_evidence(x, y, 499, 500, 19)
    second = spearman_evidence(x**3, np.exp(y), 499, 500, 19)
    assert first == second
    perfect = spearman_evidence(x, x, 499, 500, 19)
    assert perfect["spearman_rho"] == pytest.approx(1)
    assert perfect["permutation_p_value"] == pytest.approx(1 / 500)
    assert perfect["conditional_95_lower"] == pytest.approx(1)
    with pytest.raises(ValueError, match="Variable items"):
        spearman_evidence(x, np.ones(len(x)), 499, 500, 19)


def test_ordinal_wlsmv_uses_identified_threshold_tests_and_support_gate(tmp_path) -> None:
    executable = os.environ.get("MP1_R_EXECUTABLE") or shutil.which("R")
    assert executable, "Install R, lavaan, semTools and jsonlite as documented in the Colab guide."
    rng = np.random.default_rng(6951)
    n = 600
    latent = rng.normal(size=(n, 1))
    responses = 0.75 * latent + np.sqrt(1 - 0.75**2) * rng.normal(size=(n, 6))
    observed = np.digitize(responses, [-1, -0.35, 0.35, 1]) + 1
    frame = pd.DataFrame(np.vstack([observed, observed]), columns=[f"q{i}" for i in range(6)])
    frame["Country"] = np.repeat(["USA", "Indonesia"], n)
    for i in range(6):
        frame[f"d{i}"] = frame[f"q{i}"]
    frame.loc[n:, "d0"] = np.digitize(responses[:, 0], [-1.6, -0.1, 0.15, 0.3]) + 1
    frame["absent"] = 3
    data_path, request_path, result_path = [
        tmp_path / name for name in ["data.csv", "request.json", "result.json"]
    ]
    frame.to_csv(data_path, index=False)
    models = {
        name: {
            "items": [f"{prefix}{i}" for i in range(6)],
            "syntax": "F =~ " + " + ".join(f"{prefix}{i}" for i in range(6)),
        }
        for name, prefix in [("SameDistribution", "q"), ("ThresholdDIF", "d")]
    }
    models["AbsentCategory"] = {"items": ["absent", "q1", "q2"], "syntax": "F =~ absent + q1 + q2"}
    request_path.write_text(
        json.dumps(
            {
                "data_csv": str(data_path),
                "countries": ["USA", "Indonesia"],
                "minimum_cfi": 0.9,
                "maximum_rmsea": 0.08,
                "maximum_srmr": 0.08,
                "alpha": 0.05,
                "models": models,
            }
        )
    )
    script = Path(__file__).resolve().parents[1] / "scripts/ordinal_invariance.R"
    result = subprocess.run(
        [
            executable,
            "--vanilla",
            "--slave",
            "-f",
            str(script),
            "--args",
            str(request_path),
            str(result_path),
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=180,
    )
    assert result.returncode == 0, result.stderr
    evidence = json.loads(result_path.read_text())
    assert evidence["identification"] == "semTools Wu.Estabrook.2016 with std.lv"
    assert evidence["decisions"]["AbsentCategory"]["status"] == "blocked_category_support"
    assert evidence["decisions"]["SameDistribution"]["through"] == "intercepts"
    assert evidence["decisions"]["ThresholdDIF"]["status"] == "equality_not_supported"
    assert evidence["decisions"]["ThresholdDIF"]["at"] == "thresholds"
    contrasts = pd.DataFrame(evidence["comparisons"])
    same = contrasts[contrasts.construct.eq("SameDistribution")]
    assert same.added_constraint.tolist() == ["thresholds", "loadings", "intercepts"]
    assert same.df_difference.tolist() == [12, 5, 5]
    assert same.p_value.gt(0.9).all()
    assert contrasts.method.str.contains("Satorra-2000").all()
    assert not evidence["warnings"]
