"""Audit instrument compatibility before limited country-specific item comparisons."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from importlib.metadata import version
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import rankdata, spearmanr
from statsmodels.stats.multitest import multipletests
from threadpoolctl import threadpool_limits

from mp1.audit import guard_output, repository_path, sha256_file, verify_unchanged, write_reports
from mp1.ordinal import polychoric_matrix
from mp1.outcome_validation import verify_frozen
from mp1.paths import get_drive_root
from mp1.psychometrics import alpha, cfa_uls

MODELS = {
    "GAID": {
        "CP": [f"GAID{i}" for i in range(1, 4)],
        "NC": [f"GAID{i}" for i in range(4, 8)],
        "W": [f"GAID{i}" for i in range(8, 12)],
    },
    "CT": {"CO": [f"CT{i}" for i in range(1, 8)], "RS": [f"CT{i}" for i in range(8, 12)]},
    "TP": {"TP": [f"TP{i}" for i in range(1, 8)]},
}
SUBDIMENSIONS = {
    "CP": "Cognitive Preoccupation",
    "NC": "Negative Consequences",
    "W": "Withdrawal",
    "CO": "Critical Openness",
    "RS": "Reflective Skepticism",
    "TP": "IWPQ Task Performance",
}
ITEMS = [item for groups in MODELS.values() for items in groups.values() for item in items]


def inspect_workbook(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Inspect every cell and both sheets, retaining source values without workbook edits."""
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != ["Survey Data", "Codebook"]:
        raise ValueError("External workbook sheets differ from the audited release.")
    inventory, frames = [], {}
    for sheet in workbook:
        cells = list(sheet.iter_rows())
        formulas = sum(cell.data_type == "f" for row in cells for cell in row)
        if formulas or sheet.merged_cells.ranges or sheet.sheet_state != "visible":
            raise ValueError("Formula, merged, or hidden source content requires a fresh audit.")
        values = [[cell.value for cell in row] for row in cells]
        headers = values[0]
        if any(header is None for header in headers) or len(set(headers)) != len(headers):
            raise ValueError("Workbook headers must be complete and unique.")
        frames[sheet.title] = pd.DataFrame(values[1:], columns=headers)
        inventory.append(
            {
                "sheet": sheet.title,
                "rows_including_header": sheet.max_row,
                "columns": sheet.max_column,
                "nonempty_cells": sum(v is not None for r in values for v in r),
                "formula_cells": formulas,
                "comment_cells": sum(c.comment is not None for r in cells for c in r),
                "hidden_rows": sum(d.hidden for d in sheet.row_dimensions.values()),
                "hidden_columns": sum(d.hidden for d in sheet.column_dimensions.values()),
                "all_cells_inspected": True,
            }
        )
    workbook.close()
    data, codebook = frames["Survey Data"], frames["Codebook"]
    if codebook.Variable.tolist() != list(data) or codebook.Variable.duplicated().any():
        raise ValueError("Codebook must cover every data column exactly once in source order.")
    for groups in MODELS.values():
        for name, items in groups.items():
            rows = codebook.set_index("Variable").loc[items]
            if not rows.Subdimension.eq(SUBDIMENSIONS[name]).all():
                raise ValueError(
                    "Codebook subdimension membership differs from the reviewed model."
                )
            if not rows["Response coding / Values"].str.contains("1 = Strongly disagree").all():
                raise ValueError("Item response anchors require a new codebook review.")
    return data, codebook, pd.DataFrame(inventory)


def audit_external_sample(
    data: pd.DataFrame, countries: list[str]
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Reconcile the provided analytic flag, complete batteries and country identities."""
    required = ["RespondentID", "Country", "AnalyticSample", "GenAI_Frequency", *ITEMS]
    if not set(required).issubset(data):
        raise ValueError("External data are missing required source columns.")
    if data.RespondentID.isna().any() or not data.RespondentID.is_unique:
        raise ValueError("External anonymous identifiers are missing or repeated.")
    if set(data.Country) != set(countries) or not data.AnalyticSample.isin(["Yes", "No"]).all():
        raise ValueError("Unexpected country or analytic flag; no country recoding is applied.")
    values = data[ITEMS]
    if (values.notna() & ~values.isin(range(1, 6))).any().any():
        raise ValueError("External item responses must be missing or integers 1 through 5.")
    active = data.GenAI_Frequency.isin(["Sometimes", "Often", "Always"])
    if not data.GenAI_Frequency.isin(["Never", "Rarely", "Sometimes", "Often", "Always"]).all():
        raise ValueError("Unexpected GenAI frequency labels.")
    complete, absent = values.notna().all(axis=1), values.isna().all(axis=1)
    supplied = data.AnalyticSample.eq("Yes")
    if not supplied.equals(active & complete) or not absent[~supplied].all():
        raise ValueError(
            "Provided analytic flag, use eligibility and missingness do not reconcile."
        )
    rows = []
    for country in countries:
        selected = data.Country.eq(country)
        analytic = data[selected & supplied]
        rows.append(
            {
                "country": country,
                "released_records": int(selected.sum()),
                "analytic_records": len(analytic),
                "whole_battery_missing_nonactive": int((selected & absent).sum()),
                "partial_item_missing_records": int((selected & ~(complete | absent)).sum()),
                "duplicate_full_records_without_identifier": int(
                    analytic.drop(columns="RespondentID").duplicated().sum()
                ),
                "invariant_29_item_vectors": int(analytic[ITEMS].nunique(axis=1).eq(1).sum()),
                "source_rows_deleted": 0,
                "eligibility": (
                    "Source AnalyticSample=Yes; all items observed; use at least Sometimes"
                ),
            }
        )
    return data[supplied].copy(), pd.DataFrame(rows)


def spearman_evidence(
    x: np.ndarray, y: np.ndarray, permutations: int, repetitions: int, seed: int
) -> dict:
    """Use tie-aware ranks, a two-sided permutation test and paired percentile resampling."""
    x, y = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    if x.ndim != 1 or x.shape != y.shape or len(x) < 20 or not np.isfinite([x, y]).all():
        raise ValueError("Complete paired item vectors with at least 20 records are required.")
    if min(np.ptp(x), np.ptp(y)) == 0 or min(permutations, repetitions) < 100 or seed < 0:
        raise ValueError("Variable items and at least 100 permutations/resamples are required.")
    rx, ry = rankdata(x), rankdata(y)
    rx, ry = rx - rx.mean(), ry - ry.mean()
    denominator = np.sqrt(np.dot(rx, rx) * np.dot(ry, ry))
    observed = float(np.dot(rx, ry) / denominator)
    if not np.isclose(observed, spearmanr(x, y).statistic, rtol=0, atol=1e-12):
        raise ArithmeticError("Independent Spearman implementations disagree.")
    rng = np.random.default_rng(seed)
    extreme = sum(
        abs(np.dot(rx, rng.permutation(ry)) / denominator) >= abs(observed) - 1e-14
        for _ in range(permutations)
    )
    p = (1 + extreme) / (1 + permutations)
    draws = rng.integers(0, len(x), size=(repetitions, len(x)))
    a, b = rankdata(x[draws], axis=1), rankdata(y[draws], axis=1)
    a, b = a - a.mean(axis=1)[:, None], b - b.mean(axis=1)[:, None]
    denom = np.sqrt((a * a).sum(axis=1) * (b * b).sum(axis=1))
    invalid = int((denom == 0).sum())
    if invalid:
        raise ValueError("Constant bootstrap item vectors prevent the requested interval.")
    low, high = np.quantile((a * b).sum(axis=1) / denom, [0.025, 0.975])
    return {
        "n": len(x),
        "spearman_rho": observed,
        "permutation_p_value": p,
        "permutation_mc_standard_error": float(np.sqrt(p * (1 - p) / (permutations + 1))),
        "conditional_95_lower": float(low),
        "conditional_95_upper": float(high),
        "permutations": permutations,
        "bootstrap_repetitions": repetitions,
        "seed": seed,
        "invalid_bootstrap_replicates": invalid,
        "inference_scope": "Exploratory observed-item association within this country sample",
    }


def run_ordinal_engine(data: pd.DataFrame, settings: dict, private: Path, script: Path) -> dict:
    """Use R only for correctly identified ordinal WLSMV measurement models."""
    executable = os.environ.get("MP1_R_EXECUTABLE") or shutil.which("R")
    if not executable:
        raise RuntimeError("R with lavaan, semTools and jsonlite is required; see the Colab guide.")
    private.mkdir(parents=True, exist_ok=True)
    data_path = private / "external_model_input.csv"
    data[["Country", *ITEMS]].to_csv(data_path, index=False)
    request = {
        "data_csv": str(data_path.resolve()),
        "countries": settings["country_order"],
        **settings["ordinal_screen"],
        "models": {
            name: {
                "items": [i for ids in groups.values() for i in ids],
                "syntax": "\n".join(f"{g} =~ " + " + ".join(ids) for g, ids in groups.items()),
            }
            for name, groups in MODELS.items()
        },
    }
    request_path, result_path = private / "ordinal_request.json", private / "ordinal_result.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    if result_path.exists():
        result_path.unlink()
    result = subprocess.run(
        [
            executable,
            "--vanilla",
            "--slave",
            "-f",
            str(script),
            "--args",
            str(request_path),
            str(result_path),
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=300,
    )
    (private / "ordinal_engine.log").write_text(result.stdout + result.stderr, encoding="utf-8")
    if result.returncode or not result_path.is_file():
        raise RuntimeError("Ordinal measurement engine failed; inspect its private log.")
    return json.loads(result_path.read_text())


@threadpool_limits.wrap(limits=1)
def run_external_validation(
    config_path: str | Path = "config/analysis.yaml",
    settings_path: str | Path = "config/work_package_d.yaml",
    workbook_path: str | Path | None = None,
) -> Path:
    """Preserve country samples and write audited compatibility and measurement evidence."""
    config, settings, fingerprints = verify_frozen(config_path, settings_path)
    options = settings["external"]
    if workbook_path is None:
        root = get_drive_root(config["project"]["drive_root_env"])
        workbook_path = root / config["data"]["us_indonesia_root"] / options["workbook_filename"]
    workbook_path = Path(workbook_path).resolve()
    digest = sha256_file(workbook_path)
    if digest != options["workbook_sha256"]:
        raise ValueError("External workbook differs from the reviewed official release.")
    fingerprints[workbook_path] = digest
    data, codebook, inventory = inspect_workbook(workbook_path)
    original = data.copy(deep=True)
    countries = options["country_order"]
    analytic, accounting = audit_external_sample(data, countries)
    population = []
    for country in countries:
        for view, source in [("released", data), ("source_analytic", analytic)]:
            selected = source[source.Country.eq(country)]
            for field in [
                "Age_Range",
                "Gender",
                "Nationality",
                "Field_of_Study",
                "Level_of_Study",
                "Year_of_Study",
                "GenAI_Experience",
                "GenAI_Frequency",
            ]:
                for value, count in selected[field].value_counts(dropna=False).items():
                    population.append(
                        {
                            "country": country,
                            "sample": view,
                            "field": field,
                            "source_label": value,
                            "records": int(count),
                            "n": len(selected),
                            "fraction": count / len(selected),
                            "locally_recoded": False,
                        }
                    )
    crosswalk_path = Path(settings_path).resolve().with_name("work_package_d_crosswalk.csv")
    fingerprints[crosswalk_path] = sha256_file(crosswalk_path)
    crosswalk = pd.read_csv(crosswalk_path, keep_default_na=False)
    if (
        crosswalk.external_item.tolist() != ITEMS
        or crosswalk.exact_common_anchor.astype(str).ne("False").any()
    ):
        raise ValueError(
            "The reviewed crosswalk must cover all items without invented common anchors."
        )
    crosswalk = crosswalk.merge(
        codebook[["Variable", "Construct", "Subdimension"]],
        left_on="external_item",
        right_on="Variable",
        validate="one_to_one",
    )
    distributions, domains, reliabilities, factor_rows, loading_rows, phi_rows, warnings = (
        [],
        [],
        [],
        [],
        [],
        [],
        [],
    )
    for country in countries:
        selected = analytic[analytic.Country.eq(country)]
        for item in ITEMS:
            counts = selected[item].value_counts()
            domains.append(
                {
                    "country": country,
                    "item": item,
                    "n": len(selected),
                    "observed_categories": int(len(counts)),
                    "absent_categories": ";".join(str(k) for k in range(1, 6) if k not in counts),
                    "minimum_observed_category_count": int(counts.min()),
                    "count_below_five_categories": sum(counts.get(k, 0) < 5 for k in range(1, 6)),
                }
            )
            for category in range(1, 6):
                count = int(counts.get(category, 0))
                distributions.append(
                    {
                        "country": country,
                        "item": item,
                        "response": category,
                        "records": count,
                        "n": len(selected),
                        "fraction": count / len(selected),
                    }
                )
        for construct, groups in MODELS.items():
            items = [i for group in groups.values() for i in group]
            values = selected[items].to_numpy(dtype=int)
            correlation, diagnostic = polychoric_matrix(values)
            fitted = cfa_uls(correlation, groups, len(values))
            factor_rows.append(
                {
                    "country": country,
                    "construct": construct,
                    "n": len(values),
                    "estimator": "Country-specific polychoric ULS diagnostic, not invariance",
                    "proper": fitted["proper"],
                    "minimum_correlation_eigenvalue": float(np.linalg.eigvalsh(correlation).min()),
                    "boundary_polychorics": int(diagnostic.boundary.sum()),
                    "maximum_cell_probability_error": float(
                        diagnostic.maximum_cell_probability_error.max()
                    ),
                    "rmsr_off_diagonal": fitted["rmsr_off_diagonal"],
                    "maximum_absolute_residual": fitted["maximum_absolute_residual"],
                    "maximum_factor_correlation": fitted["maximum_factor_correlation"],
                }
            )
            loading_rows.extend(
                {"country": country, "construct": construct, **row}
                for row in fitted["loadings"].to_dict("records")
            )
            phi = fitted["factor_correlations"]
            for i in range(len(phi)):
                for j in range(i):
                    phi_rows.append(
                        {
                            "country": country,
                            "construct": construct,
                            "factor_1": phi.index[i],
                            "factor_2": phi.index[j],
                            "correlation": float(phi.iloc[i, j]),
                        }
                    )
            warnings.extend(
                {"country": country, "construct": construct, "message": w}
                for w in fitted["warnings"]
            )
            for name, members in groups.items():
                reliabilities.append(
                    {
                        "country": country,
                        "subdimension": SUBDIMENSIONS[name],
                        "items": ";".join(members),
                        "n": len(selected),
                        "raw_alpha": alpha(np.cov(selected[members].to_numpy(), rowvar=False)),
                        "validated_scale_claim": False,
                    }
                )
    private = repository_path(config_path, settings["d3"]["private_directory"])
    if not private.is_relative_to(repository_path(config_path, "data/interim")):
        raise ValueError("Private external files must remain under ignored data/interim.")
    script = repository_path(config_path, "scripts/ordinal_invariance.R")
    fingerprints[script] = sha256_file(script)
    for name in [
        "external_model_input.csv",
        "ordinal_request.json",
        "ordinal_result.json",
        "ordinal_engine.log",
    ]:
        guard_output(private / name, fingerprints)
    engine = run_ordinal_engine(analytic, options, private, script)
    associations = []
    for c, country in enumerate(countries):
        selected = analytic[analytic.Country.eq(country)]
        for j, (item_x, item_y) in enumerate(options["observed_item_relationships"]):
            result = spearman_evidence(
                selected[item_x].to_numpy(),
                selected[item_y].to_numpy(),
                options["permutation_repetitions"],
                options["bootstrap_repetitions"],
                options["seed"] + 10 * c + j,
            )
            associations.append({"country": country, "item_x": item_x, "item_y": item_y, **result})
    associations = pd.DataFrame(associations)
    associations["holm_p_four_test_family"] = multipletests(
        associations.permutation_p_value, method="holm"
    )[1]
    bangladesh = pd.read_csv(
        repository_path(config_path, "outputs/psychometrics/item_inventory.csv")
    )
    coverage = []
    for row in bangladesh.itertuples():
        matches = crosswalk[
            crosswalk[f"{row.dataset}_related_items"]
            .str.split(";")
            .map(lambda x, item=row.item: item in x)
        ]
        coverage.append(
            {
                "dataset": row.dataset,
                "item": row.item,
                "related_external_items": ";".join(matches.external_item),
                "relationship": "theme_only" if len(matches) else "no_item_counterpart",
                "direct_anchor": False,
                "score_transfer_authorized": False,
            }
        )
    fields = []
    for row in codebook.itertuples(index=False, name=None):
        name, _, construct, subdimension, _, _ = row
        fields.append(
            {
                "field": name,
                "construct": construct,
                "subdimension": subdimension
                if subdimension in SUBDIMENSIONS.values()
                else "not_applicable",
                "source_nonmissing": int(data[name].notna().sum()),
                "source_distinct_values": int(data[name].nunique()),
                "country_identifier": name == "Country",
                "used_in_item_analysis": name in ITEMS,
                "codebook_row_inspected": True,
            }
        )
    table_list = {
        "workbook_inventory.csv": inventory,
        "field_audit.csv": pd.DataFrame(fields),
        "sample_accounting.csv": accounting,
        "population_marginals.csv": pd.DataFrame(population),
        "item_compatibility.csv": crosswalk.drop(columns="Variable"),
        "bangladesh_item_coverage.csv": pd.DataFrame(coverage),
        "item_distributions.csv": pd.DataFrame(distributions),
        "item_category_support.csv": pd.DataFrame(domains),
        "subdimension_reliability.csv": pd.DataFrame(reliabilities),
        "country_factor_diagnostics.csv": pd.DataFrame(factor_rows),
        "country_factor_loadings.csv": pd.DataFrame(loading_rows),
        "country_factor_correlations.csv": pd.DataFrame(phi_rows),
        "ordinal_model_fit.csv": pd.DataFrame(engine["models"]),
        "ordinal_model_parameters.csv": pd.DataFrame(engine["parameters"]),
        "ordinal_equality_comparisons.csv": pd.DataFrame(
            engine["comparisons"],
            columns=[
                "construct",
                "added_constraint",
                "method",
                "chi_square_difference",
                "df_difference",
                "p_value",
            ],
        ),
        "observed_item_associations.csv": associations,
    }
    if not data.equals(original):
        raise RuntimeError("External source values or order changed.")
    summary = {
        "package": "D",
        "baseline_commit": settings["baseline_commit"],
        "python_hash_seed": os.environ.get("PYTHONHASHSEED", "not_set_use_cli_for_reproduction"),
        "source_sha256": digest,
        "source_url": options["source_url"],
        "settings_sha256": sha256_file(settings_path),
        "crosswalk_sha256": sha256_file(crosswalk_path),
        "code_sha256": {
            "cross_cultural.py": sha256_file(__file__),
            "ordinal_invariance.R": sha256_file(script),
        },
        "versions": {
            name: version(name) for name in ["numpy", "pandas", "scipy", "openpyxl", "semopy"]
        },
        "ordinal_engine": {
            k: v for k, v in engine.items() if k not in ["parameters", "models", "comparisons"]
        },
        "warnings": warnings,
        "bangladesh_compatibility": (
            "No identical anchors; content, response meaning and populations differ. No "
            "pooling, score harmonization, profile transfer or latent invariance test "
            "with Bangladesh."
        ),
        "cross_country_method": (
            "Parallel country-specific observed-item Spearman associations with tie-aware "
            "permutation tests and Holm correction. No latent means, total scores, "
            "standardized country scores or country-causal tests."
        ),
        "relationships": options["observed_item_relationships"],
        "relationship_scope": (
            "GAID6 concerns confidence without AI; CT9 concerns general credibility "
            "checking; TP6 concerns self-rated efficiency. These are individual "
            "responses, not validated scales or CGPA."
        ),
        "population": (
            "Separate U.S. and Indonesian undergraduate business-related samples; no "
            "representative national sampling claim. Country is the recorded national "
            "sample, not a randomized culture exposure."
        ),
        "translation_evidence": (
            "A common English Codebook is supplied; original administered Indonesian item "
            "forms and translation-validation evidence are absent from this release."
        ),
        "performance_conflict": (
            "Advisor guidance suggesting more objective performance is unsupported: "
            "TP1-TP7 are self-reported agreement items, not recorded grades or a "
            "performance test."
        ),
        "standardization_decision": (
            "Within-country z-scores cannot create common item content or measurement "
            "equivalence and would erase country mean differences by construction; not "
            "used."
        ),
        "invariance_interpretation": (
            "A failed configural/support screen is not a formal rejection of every "
            "possible invariant model. No partial-invariance, item-deletion or "
            "category-merging search is performed."
        ),
        "raw_rows_deleted": 0,
        "source_values_and_order_unchanged": True,
        "country_identities_preserved": True,
        "vietnam_used": False,
        "acceptance_status": "requires_manual_scientific_review",
    }
    destination = repository_path(config_path, options["output_directory"])
    write_reports(destination, table_list, summary, "external_summary.json", fingerprints)
    verify_unchanged(fingerprints)
    verify_frozen(config_path, settings_path)
    return destination
